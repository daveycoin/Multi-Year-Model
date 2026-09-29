"""Builds the 5-year forecast + cash flow workbook driven by a QuickBooks budget export.

Usage: python3 build_model.py <quickbooks_export.xlsx> <output.xlsx>
"""
import sys
import datetime as dt
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule

SRC, OUT = sys.argv[1], sys.argv[2]
NY = 5  # forecast years

# ---------------------------------------------------------------- styles
FONT = "Arial"
f_norm = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_in = Font(name=FONT, size=10, color="0000FF")
f_title = Font(name=FONT, size=14, bold=True, color="1F3864")
f_sub = Font(name=FONT, size=10, italic=True, color="595959")
f_hdr = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_link = Font(name=FONT, size=10, color="008000")
fill_hdr = PatternFill("solid", fgColor="1F3864")
fill_sec = PatternFill("solid", fgColor="D9E1F2")
fill_in = PatternFill("solid", fgColor="FFF2CC")
fill_tot = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="808080")
top = Border(top=thin)
topbot = Border(top=thin, bottom=Side(style="double", color="808080"))
NUM = '#,##0;(#,##0);"-"'
PCT = '0.0%'
DATE = 'mmm-yy'


def put(ws, ref, v, font=f_norm, fmt=None, fill=None, al=None, border=None):
    c = ws[ref]
    c.value = v
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if al:
        c.alignment = al
    if border:
        c.border = border
    return c


def inp(ws, ref, v=None, fmt=None):
    return put(ws, ref, v, f_in, fmt, fill_in)


def hdr_row(ws, row, c1, c2):
    for c in range(c1, c2 + 1):
        x = ws.cell(row=row, column=c)
        x.font, x.fill = f_hdr, fill_hdr
        x.alignment = Alignment(horizontal="center" if c > c1 else "left", vertical="center", wrap_text=True)


def section(ws, row, text, c2):
    for c in range(1, c2 + 1):
        ws.cell(row=row, column=c).fill = fill_sec
    put(ws, f"A{row}", text, f_bold, fill=fill_sec)


# ---------------------------------------------------------------- parse QB export
src_wb = openpyxl.load_workbook(SRC)
src = src_wb.active
SECT = {"Revenue": "Revenue", "Cost of Goods Sold": "COGS", "Expenditures": "Expense",
        "Other Revenue": "Other Revenue", "Other Expenditures": "Other Expense"}
NET_ROWS = {"Gross Profit", "Net Operating Revenue", "Net Other Revenue", "Net Revenue"}
lines, section_now, cat_now = [], "", ""
for r in range(1, src.max_row + 1):
    lab = src.cell(r, 1).value
    lab = (lab or "").strip() if isinstance(lab, str) else ("" if lab is None else str(lab))
    val = src.cell(r, 2).value
    if lab in SECT:
        section_now, cat_now = SECT[lab], SECT[lab]
        continue
    if len(lab) > 2 and lab[1] == " " and lab[0].isdigit():
        cat_now = lab
    is_num = isinstance(val, (int, float)) or (isinstance(val, str) and val.startswith("="))
    if lab and is_num and not lab.startswith("Total for") and lab not in NET_ROWS:
        acct = lab[:6] if len(lab) > 7 and lab[:6].isdigit() and lab[6] == " " else None
        key = acct or lab
        base = val if isinstance(val, (int, float)) else 0
        lines.append(dict(type="Revenue" if "Revenue" in section_now else "Expense",
                          cat=cat_now, key=key, desc=lab, base=base))
keys = [l["key"] for l in lines]
assert len(keys) == len(set(keys)), "duplicate keys in export"

# ---------------------------------------------------------------- default drivers
MAN = {  # key -> manual FY28..FY31
    "400187": [380000, 360000, 360000, 360000],
    "400180": [0, 0, 0, 0],
    "400126": [3000, 3000, 3000, 3000],
    "400194": [-325000, -325000, -325000, -325000],
    "400191": [50000, 0, 0, 0],
}
METHOD = {
    "400186": "Tuition Index", "400188": "Flat", "400181": "Enrollment", "400183": "Enrollment",
    "400124": "Reserve Yield", "500267": "% of Salaries", "500274": "% of Salaries",
    "500273": "% of Salaries", "500268": "Benefits", "500269": "Benefits", "500270": "Tuition Index",
    "500272": "Flat", "500275": "Salary", "500276": "Salary", "500277": "Salary", "500278": "Salary",
    "500279": "Salary", "500280": "Salary", "Food Service": "Per Student",
    "500123": "Debt", "500125": "Debt", "500393": "Enrollment",
}
FLAT_REV = {"400193", "400195", "400196", "400133", "400134", "400197", "400198", "400080", "400081",
            "400084", "400192"}
TUITION_PLAN = {"400186", "400187", "400188", "400189", "400193", "400194", "400195", "400196", "400183"}
JUL_LUMP = {"400180", "400181", "400184", "400185", "400199", "400201", "400080", "400081", "400084",
            "400191", "400192"}
for l in lines:
    k = l["key"]
    if k in MAN:
        l["method"] = "Manual"
    elif k in METHOD:
        l["method"] = METHOD[k]
    elif l["type"] == "Revenue" or l["base"] == 0:
        l["method"] = "Flat"
    else:
        l["method"] = "Inflation"
    l["profile"] = "Tuition Plan" if k in TUITION_PLAN else ("Jul Lump" if k in JUL_LUMP else "Even")
    l["manual"] = MAN.get(k)

METHODS = ["Flat", "Inflation", "Salary", "Benefits", "Per Student", "Enrollment", "Tuition Index",
           "% of Salaries", "Debt", "Reserve Yield", "Manual"]
PROFILES = ["Tuition Plan", "Even", "Jul Lump", "Custom A", "Custom B"]

wb = openpyxl.Workbook()
wsS = wb.active
wsS.title = "Summary"
wsI = wb.create_sheet("Inputs")
wsE = wb.create_sheet("Enrollment")
wsF = wb.create_sheet("Forecast Detail")
wsC = wb.create_sheet("Cash Flow")
wsCS = wb.create_sheet("CF Summary")
wsQ = wb.create_sheet("QB Import")
wsR = wb.create_sheet("Read Me")
for ws in wb:
    ws.sheet_view.showGridLines = False
YC = ["C", "D", "E", "F", "G"]          # Inputs year columns
FY = ["G", "H", "I", "J", "K"]          # Forecast Detail year columns
SC = ["B", "C", "D", "E", "F"]          # Summary / Enrollment / CF Summary year columns

# ================================================================= QB IMPORT
QN = 600
for r in range(1, src.max_row + 1):
    for c in (1, 2, 3):
        v = src.cell(r, c).value
        if v is not None:
            wsQ.cell(r, c).value = v
            wsQ.cell(r, c).font = f_norm
            if c > 1:
                wsQ.cell(r, c).number_format = NUM
for r in range(1, QN + 1):
    prevE = f"E{r-1}" if r > 1 else '""'
    prevF = f"F{r-1}" if r > 1 else '""'
    put(wsQ, f"D{r}", f'=TRIM(A{r})', f_link)
    put(wsQ, f"E{r}", f'=IF(D{r}="Revenue","Revenue",IF(D{r}="Cost of Goods Sold","COGS",IF(D{r}="Expenditures","Expense",'
        f'IF(D{r}="Other Revenue","Other Revenue",IF(D{r}="Other Expenditures","Other Expense",{prevE})))))', f_link)
    put(wsQ, f"F{r}", f'=IF(OR(D{r}="Revenue",D{r}="Cost of Goods Sold",D{r}="Expenditures",D{r}="Other Revenue",D{r}="Other Expenditures"),E{r},'
        f'IF(AND(LEN(D{r})>2,MID(D{r},2,1)=" ",ISNUMBER(VALUE(LEFT(D{r},1)))),D{r},{prevF}))', f_link)
    put(wsQ, f"G{r}", f'=IF(AND(D{r}<>"",ISNUMBER(B{r}),LEFT(D{r},9)<>"Total for",D{r}<>"Gross Profit",D{r}<>"Net Operating Revenue",'
        f'D{r}<>"Net Other Revenue",D{r}<>"Net Revenue"),1,0)', f_link)
    put(wsQ, f"H{r}", f'=IF(G{r}=1,IF(AND(LEN(D{r})>7,ISNUMBER(VALUE(LEFT(D{r},6))),MID(D{r},7,1)=" "),LEFT(D{r},6),D{r}),"")', f_link)
    put(wsQ, f"I{r}", f"=IF(G{r}=1,IF(COUNTIF('Forecast Detail'!$C$8:$C$400,H{r})=0,\"UNMAPPED\",\"OK\"),\"\")", f_link)
for c, w in zip("ABCDEFGHI", [52, 16, 14, 4, 16, 34, 9, 40, 12]):
    wsQ.column_dimensions[c].width = w
# instructions box to the right
put(wsQ, "K1", "HOW TO IMPORT", f_bold)
for i, t in enumerate(["1. In QuickBooks, export the Budget Overview P&L report to Excel.",
                       "2. Clear columns A:C on this tab (do NOT touch columns D:I).",
                       "3. Paste the whole export so its first row lands in A1.",
                       "4. Check the status line at the top of the Summary tab.",
                       "Columns D:I (green) are helper formulas that read the pasted data.",
                       "Any account flagged UNMAPPED in column I is held flat in the forecast",
                       "under its Major Category; add it to Forecast Detail to forecast it."], start=2):
    put(wsQ, f"K{i}", t)
wsQ.column_dimensions["K"].width = 70
wsQ.freeze_panes = "A2"

# ================================================================= INPUTS
ws = wsI
put(ws, "A1", "INPUTS & ASSUMPTIONS", f_title)
put(ws, "A2", "Blue-on-yellow cells are inputs. Everything else is calculated. FY27 amounts come from the QuickBooks import.", f_sub)
ws.column_dimensions["A"].width = 46
for c in range(2, 17):
    ws.column_dimensions[L(c)].width = 12.5
R_LBL, R_END, R_START = 4, 5, 6
put(ws, f"A{R_LBL}", "Fiscal year", f_hdr)
hdr_row(ws, R_LBL, 1, 7)
put(ws, f"A{R_END}", "Fiscal year ending")
put(ws, f"A{R_START}", "Fiscal year starting")
for i, c in enumerate(YC):
    if i == 0:
        inp(ws, f"{c}{R_END}", dt.date(2027, 6, 30), "m/d/yyyy")
    else:
        put(ws, f"{c}{R_END}", f"=EOMONTH({YC[i-1]}{R_END},12)", fmt="m/d/yyyy")
    put(ws, f"{c}{R_LBL}", f'="FY"&TEXT({c}{R_END},"yy")', f_hdr, fill=fill_hdr, al=Alignment(horizontal="center"))
    put(ws, f"{c}{R_START}", f"=EOMONTH({c}{R_END},-12)+1", fmt="m/d/yyyy")
R_TUIT, R_SAL, R_BEN, R_INFL, R_YLD, R_DEBTEND, R_DEBTMO, R_INTBASE = 9, 10, 11, 12, 13, 14, 15, 16
section(ws, 8, "Growth assumptions", 7)
rows = [(R_TUIT, "Tuition rate increase %", [0.04, 0.05, 0.045, 0.045, 0.04], "FY27 shown for reference; rates already set in the budget"),
        (R_SAL, "Salary increase %", [0.025, 0.03, 0.03, 0.03, 0.03], "Applies to lines using the 'Salary' method"),
        (R_BEN, "Benefits cost increase % (medical, dental)", [0.03, 0.03, 0.03, 0.03, 0.03], "Lines using the 'Benefits' method"),
        (R_INFL, "General inflation %", [0.03, 0.03, 0.03, 0.03, 0.03], "Lines using 'Inflation' and 'Per Student'"),
        (R_YLD, "Yield on interest-earning balances %", [None, 0.014, 0.014, 0.014, 0.014], "Applied to prior year-end balances; 'Reserve Yield' method")]
for r, lab, vals, note in rows:
    put(ws, f"A{r}", lab)
    for c, v in zip(YC, vals):
        if v is not None:
            inp(ws, f"{c}{r}", v, PCT)
    put(ws, f"H{r}", note, f_sub)
put(ws, f"A{R_DEBTEND}", "Final mortgage payment date")
inp(ws, f"C{R_DEBTEND}", dt.date(2029, 10, 31), "m/d/yyyy")
put(ws, f"H{R_DEBTEND}", "Lines using the 'Debt' method stop after this month", f_sub)
put(ws, f"A{R_DEBTMO}", "Months of debt service paid in fiscal year")
for c in YC:
    put(ws, f"{c}{R_DEBTMO}", f"=MAX(0,MIN(12,(YEAR($C${R_DEBTEND})-YEAR({c}{R_START}))*12+MONTH($C${R_DEBTEND})-MONTH({c}{R_START})+1))", fmt="0")
put(ws, f"A{R_INTBASE}", "Interest-earning balance at prior year-end")

# --- strategic plan initiatives
R_IT = 18
section(ws, R_IT, "Strategic plan initiatives  (Operating = flows through the P&L; Capital = cash only, no P&L)", 14)
h = ["Initiative", "Type", "On? (Y/N)", "Spend month (0=even, 1=Jul)"] + \
    [f'="Cost "&{c}${R_LBL}' for c in YC] + [f'="Offset "&{c}${R_LBL}' for c in YC]
for i, t in enumerate(h):
    put(ws, f"{L(i+1)}{R_IT+1}", t)
hdr_row(ws, R_IT + 1, 1, 14)
ws.row_dimensions[R_IT + 1].height = 40
IT1, IT2 = R_IT + 2, R_IT + 11      # rows 20..29
seed = {0: ("Additional Lower School teacher (per prior model: +$50k/yr)", "Operating", "Y", 0, [0, 50000, 100000, 150000, 200000], [0] * 5),
        1: ("Lower School reconfiguration (enter cost)", "Capital", "Y", 0, [0, 0, 0, 0, 0], [0] * 5)}
for i, r in enumerate(range(IT1, IT2 + 1)):
    s = seed.get(i)
    inp(ws, f"A{r}", s[0] if s else None)
    inp(ws, f"B{r}", s[1] if s else None)
    inp(ws, f"C{r}", s[2] if s else None)
    inp(ws, f"D{r}", s[3] if s else None, "0")
    for j in range(5):
        inp(ws, f"{L(5+j)}{r}", s[4][j] if s else None, NUM)
        inp(ws, f"{L(10+j)}{r}", s[5][j] if s else None, NUM)
R_ITOP, R_ITCAP = IT2 + 2, IT2 + 3   # 31, 32
put(ws, f"A{R_ITOP}", "Operating initiatives - net cost to P&L", f_bold)
put(ws, f"A{R_ITCAP}", "Capital initiatives - net cash cost", f_bold)
for j in range(5):
    c, o = L(5 + j), L(10 + j)
    put(ws, f"{c}{R_ITOP}", f'=SUMPRODUCT(($B${IT1}:$B${IT2}="Operating")*($C${IT1}:$C${IT2}="Y")*({c}${IT1}:{c}${IT2}-{o}${IT1}:{o}${IT2}))', f_bold, NUM, fill_tot, border=top)
    put(ws, f"{c}{R_ITCAP}", f'=SUMPRODUCT(($B${IT1}:$B${IT2}="Capital")*($C${IT1}:$C${IT2}="Y")*({c}${IT1}:{c}${IT2}-{o}${IT1}:{o}${IT2}))', f_bold, NUM, fill_tot)
dv = DataValidation(type="list", formula1='"Operating,Capital"', allow_blank=True)
dv2 = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
ws.add_data_validation(dv); ws.add_data_validation(dv2)
dv.add(f"B{IT1}:B{IT2}"); dv2.add(f"C{IT1}:C{IT2}")

# --- capital projects
R_CP = R_ITCAP + 2   # 34
section(ws, R_CP, "Capital projects & equipment (cash only; outside the operating P&L)", 7)
put(ws, f"A{R_CP+1}", "Project"); put(ws, f"B{R_CP+1}", "Spend month (0=even, 1=Jul)")
for c in YC:
    put(ws, f"{c}{R_CP+1}", f"={c}${R_LBL}")
hdr_row(ws, R_CP + 1, 1, 7)
ws.row_dimensions[R_CP + 1].height = 40
CP1, CP2 = R_CP + 2, R_CP + 7
for r in range(CP1, CP2 + 1):
    inp(ws, f"A{r}"); inp(ws, f"B{r}", 0 if r > CP1 else 0, "0")
    for c in YC:
        inp(ws, f"{c}{r}", None, NUM)
ws[f"A{CP1}"].value = "Example: roof replacement"
R_CPT = CP2 + 1
put(ws, f"A{R_CPT}", "Total capital projects", f_bold)
for c in YC:
    put(ws, f"{c}{R_CPT}", f"=SUM({c}{CP1}:{c}{CP2})", f_bold, NUM, fill_tot, border=top)

# --- reserve funding
R_RF = R_CPT + 2
section(ws, R_RF, "Reserve funding (transfer from operating cash to reserves; cash-neutral overall)", 7)
put(ws, f"A{R_RF+1}", "Transfer"); put(ws, f"B{R_RF+1}", "Month (0=even, 1=Jul)")
for c in YC:
    put(ws, f"{c}{R_RF+1}", f"={c}${R_LBL}")
hdr_row(ws, R_RF + 1, 1, 7)
ws.row_dimensions[R_RF + 1].height = 40
RF1, RF2 = R_RF + 2, R_RF + 3
for r, t in [(RF1, "Operating reserve funding"), (RF2, "Repair & replacement (R&R) reserve funding")]:
    put(ws, f"A{r}", t); inp(ws, f"B{r}", 12, "0")
    for c in YC:
        inp(ws, f"{c}{r}", 0, NUM)
R_RFT = RF2 + 1
put(ws, f"A{R_RFT}", "Total reserve funding", f_bold)
for c in YC:
    put(ws, f"{c}{R_RFT}", f"=SUM({c}{RF1}:{c}{RF2})", f_bold, NUM, fill_tot, border=top)

# --- opening balances
R_OB = R_RFT + 2
section(ws, R_OB, "Opening cash & reserve balances", 7)
put(ws, f"A{R_OB+1}", "Balances as of (start of FY27 = 6/30/2026 recommended)")
inp(ws, f"B{R_OB+1}", dt.date(2026, 6, 30), "m/d/yyyy")
for i, t in enumerate(["Account", "Available for operations? (Y/N)", "Earns interest? (Y/N)", "Balance"]):
    put(ws, f"{L(i+1)}{R_OB+2}", t)
hdr_row(ws, R_OB + 2, 1, 4)
ws.row_dimensions[R_OB + 2].height = 40
acc = [("Operating Account", "Y", "N", 134596.05), ("Contingency", "N", "Y", 441027.78), ("S4K", "N", "Y", 167658.20),
       ("PPP Reserve Account", "Y", "Y", 1029181.76), ("Development Account", "N", "Y", 433920.19)]
OB1 = R_OB + 3
for i, (n, a, e, b) in enumerate(acc):
    r = OB1 + i
    inp(ws, f"A{r}", n); inp(ws, f"B{r}", a); inp(ws, f"C{r}", e); inp(ws, f"D{r}", b, NUM)
OB2 = OB1 + len(acc) - 1
dv2.add(f"B{OB1}:C{OB2}")
R_OBA, R_OBR, R_OBE = OB2 + 1, OB2 + 2, OB2 + 3
put(ws, f"A{R_OBA}", "Operating cash available (opening)", f_bold)
put(ws, f"D{R_OBA}", f'=SUMIF(B{OB1}:B{OB2},"Y",D{OB1}:D{OB2})', f_bold, NUM, fill_tot, border=top)
put(ws, f"A{R_OBR}", "Reserves & restricted (opening)", f_bold)
put(ws, f"D{R_OBR}", f'=SUMIF(B{OB1}:B{OB2},"N",D{OB1}:D{OB2})', f_bold, NUM, fill_tot)
put(ws, f"A{R_OBE}", "Interest-earning balances (opening)", f_bold)
put(ws, f"D{R_OBE}", f'=SUMIF(C{OB1}:C{OB2},"Y",D{OB1}:D{OB2})', f_bold, NUM, fill_tot)
put(ws, f"E{OB1}", "Balances seeded from the 1/31/26 figures in the old model - UPDATE.", f_sub)
# interest base row
for i, c in enumerate(YC):
    if i == 0:
        continue
    put(ws, f"{c}{R_INTBASE}", f"=$D${R_OBE}+SUM($C${R_RFT}:{YC[i-1]}{R_RFT})", fmt=NUM)
put(ws, f"H{R_INTBASE}", "Opening earning balances + reserve funding to date", f_sub)

# --- payment plans
R_PP = R_OBE + 2
section(ws, R_PP, "Tuition payment plans (drives monthly cash receipts)", 16)
for i, t in enumerate(["Plan", "Families", "% of tuition"]):
    put(ws, f"{L(i+1)}{R_PP+1}", t)
hdr_row(ws, R_PP + 1, 1, 3)
PP1 = R_PP + 2
plans = [("Plan A - 1 annual payment (June 30)", 94), ("Plan B - 4 payments (May/Aug/Nov/Feb)", 96),
         ("Plan C - 12 monthly payments (May-April), incl. S4K & employees", 385)]
for i, (n, cnt) in enumerate(plans):
    r = PP1 + i
    put(ws, f"A{r}", n); inp(ws, f"B{r}", cnt, "0")
    put(ws, f"C{r}", f"=B{r}/SUM($B${PP1}:$B${PP1+2})", fmt=PCT)
R_PT = PP1 + 3
put(ws, f"A{R_PT}", "Total families in plans", f_bold)
put(ws, f"B{R_PT}", f"=SUM(B{PP1}:B{PP1+2})", f_bold, "0", fill_tot, border=top)
R_SCH = R_PT + 2
put(ws, f"A{R_SCH}", "Payment schedule: % of a plan's annual tuition received in each period", f_bold)
periods = ["Prior May", "Prior Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar", "Apr", "May", "Jun"]
def period_hdr(row, first):
    put(ws, f"A{row}", first)
    for i, p in enumerate(periods):
        put(ws, f"{L(2+i)}{row}", p)
    put(ws, f"P{row}", "Total")
    hdr_row(ws, row, 1, 16)
period_hdr(R_SCH + 1, "Plan")
SCH1 = R_SCH + 2
sched = {0: {1: 1.0},
         1: {0: .25, 3: .25, 5: .25, 8: .25},   # index into 0-based period list
         2: {0: 1 / 12, 1: 1 / 12, **{i: 1 / 12 for i in range(2, 12)}}}
sched[1] = {0: .25, 3: .25, 6: .25, 9: .25}      # May(prior), Aug, Nov, Feb
for i in range(3):
    r = SCH1 + i
    put(ws, f"A{r}", plans[i][0].split(" - ")[0])
    for p in range(14):
        v = sched[i].get(p, 0)
        inp(ws, f"{L(2+p)}{r}", v, PCT)
    put(ws, f"P{r}", f"=SUM(B{r}:O{r})", fmt=PCT)
R_PR = SCH1 + 4
put(ws, f"A{R_PR-1}", "Cash timing profiles: % of a fiscal year's amount received/paid in each period (used by Forecast Detail column F)", f_bold)
period_hdr(R_PR, "Profile")
PR1 = R_PR + 1
prof_rows = {}
for i, p in enumerate(PROFILES):
    r = PR1 + i
    prof_rows[p] = r
    put(ws, f"A{r}", p)
    if p == "Tuition Plan":
        for k in range(14):
            col = L(2 + k)
            put(ws, f"{col}{r}", f"=SUMPRODUCT($C${PP1}:$C${PP1+2},{col}{SCH1}:{col}{SCH1+2})", fmt=PCT)
    else:
        vals = [0] * 14
        if p == "Even":
            vals = [0, 0] + [1 / 12] * 12
        elif p == "Jul Lump":
            vals[2] = 1
        elif p == "Custom A":
            vals[5] = .5; vals[10] = .5
        elif p == "Custom B":
            vals[13] = 1
        for k in range(14):
            inp(ws, f"{L(2+k)}{r}", vals[k], PCT)
    put(ws, f"P{r}", f"=SUM(B{r}:O{r})", fmt=PCT)
R_PRCHK = PR1 + 5
put(ws, f"A{R_PRCHK}", "Profile rows must total 100%", f_sub)
ws.conditional_formatting.add(f"P{PR1}:P{PR1+4}", FormulaRule(formula=[f"ABS(P{PR1}-1)>0.0001"], fill=PatternFill("solid", bgColor="FFC7CE")))
ws.conditional_formatting.add(f"P{SCH1}:P{SCH1+2}", FormulaRule(formula=[f"ABS(P{SCH1}-1)>0.0001"], fill=PatternFill("solid", bgColor="FFC7CE")))
ws.freeze_panes = "B7"

# ================================================================= ENROLLMENT
ws = wsE
put(ws, "A1", "ENROLLMENT & TUITION RATES", f_title)
put(ws, "A2", "FY27 enrollment and rates are inputs. Later years roll each grade forward: prior grade x retention + new students.", f_sub)
ws.column_dimensions["A"].width = 44
for c in "BCDEFG":
    ws.column_dimensions[c].width = 13
grades = ["K-3 (5 day)", "K-3 (3 day)", "K-4", "K5", "1st", "2nd", "3rd", "4th", "5th", "6th", "7th", "8th", "9th", "10th", "11th", "12th"]
enr27 = [21, 4, 44, 53, 45, 37, 33, 36, 36, 40, 34, 37, 45, 44, 43, 30]
rate27 = [8958.56, 5400, 10510.24, 13198.64, 15895.36, 15895.36, 15895.36, 16672.24, 16672.24, 17899.44, 17899.44, 17899.44, 18075.2, 18075.2, 18075.2, 18075.2]
new = [20, 4, 13, 3, 2, 1, 1, 1, 1, 4, 2, 2, 4, 1, 1, 0]
E_H, E1 = 4, 5
E_TOT = E1 + 16      # 21
R_EG, R_EI, R_EA = E_TOT + 2, E_TOT + 3, E_TOT + 4       # 23,24,25
RT_H, RT1 = 28, 29
AS_H, AS1 = 48, 49
def yr_hdr(row, first):
    put(ws, f"A{row}", first)
    for i, c in enumerate(SC):
        put(ws, f"{c}{row}", f"=Inputs!{YC[i]}${R_LBL}")
    hdr_row(ws, row, 1, 6)
put(ws, f"A{E_H-1}", "Enrollment by grade", f_bold)
yr_hdr(E_H, "Grade")
for gi, g in enumerate(grades):
    r = E1 + gi
    ar = AS1 + gi
    put(ws, f"A{r}", g)
    inp(ws, f"B{r}", enr27[gi], "#,##0")
    for i in range(1, 5):
        X, P = SC[i], SC[i - 1]
        if gi in (0, 1):
            f = f"={X}{ar}"
        elif gi == 2:
            f = f"=ROUND(({P}{E1}+{P}{E1+1})*$B{ar}+{X}{ar},0)"
        else:
            f = f"=ROUND({P}{r-1}*$B{ar}+{X}{ar},0)"
        put(ws, f"{X}{r}", f, fmt="#,##0")
put(ws, f"A{E_TOT}", "Total enrollment", f_bold)
put(ws, f"A{E_TOT+1}", "  Change vs prior year")
put(ws, f"A{R_EG}", "Tuition potential (enrollment x rate)")
put(ws, f"A{R_EI}", "Tuition index (FY27 = 1.000)", f_bold)
put(ws, f"A{R_EA}", "Average rate per student")
for i, c in enumerate(SC):
    put(ws, f"{c}{E_TOT}", f"=SUM({c}{E1}:{c}{E1+15})", f_bold, "#,##0", fill_tot, border=top)
    if i:
        put(ws, f"{c}{E_TOT+1}", f"={c}{E_TOT}-{SC[i-1]}{E_TOT}", fmt='#,##0;(#,##0);"-"')
    put(ws, f"{c}{R_EG}", f"=SUMPRODUCT({c}{E1}:{c}{E1+15},{c}{RT1}:{c}{RT1+15})", fmt=NUM)
    put(ws, f"{c}{R_EI}", f"={c}{R_EG}/$B${R_EG}", f_bold, "0.0000")
    put(ws, f"{c}{R_EA}", f"={c}{R_EG}/{c}{E_TOT}", fmt=NUM)
put(ws, f"A{RT_H-1}", "Tuition rate by grade (before discounts)", f_bold)
yr_hdr(RT_H, "Grade")
for gi, g in enumerate(grades):
    r = RT1 + gi
    put(ws, f"A{r}", g + (" (rate not in old model - enter)" if gi == 1 else ""))
    inp(ws, f"B{r}", rate27[gi], NUM)
    for i in range(1, 5):
        put(ws, f"{SC[i]}{r}", f"={SC[i-1]}{r}*(1+Inputs!{YC[i]}${R_TUIT})", fmt=NUM)
put(ws, f"A{AS_H-1}", "Roll-forward assumptions", f_bold)
put(ws, f"A{AS_H}", "Grade")
put(ws, f"B{AS_H}", "Retention into grade")
for i in range(1, 5):
    put(ws, f"{SC[i]}{AS_H}", f'="New students "&Inputs!{YC[i]}${R_LBL}')
hdr_row(ws, AS_H, 1, 6)
ws.row_dimensions[AS_H].height = 40
for gi, g in enumerate(grades):
    r = AS1 + gi
    put(ws, f"A{r}", g)
    if gi >= 2:
        inp(ws, f"B{r}", 0.96, PCT)
    else:
        put(ws, f"B{r}", "n/a", f_sub, al=Alignment(horizontal="right"))
    for i in range(1, 5):
        inp(ws, f"{SC[i]}{r}", new[gi], "0")
put(ws, f"A{AS1+17}", "K-4 is fed by both K-3 programs. Retention is the share of the prior grade that advances; new students are additions.", f_sub)
ws.freeze_panes = "B5"

# ================================================================= FORECAST DETAIL
ws = wsF
put(ws, "A1", "FORECAST DETAIL - by QuickBooks account", f_title)
put(ws, "A2", "FY27 = QuickBooks budget import. Choose a forecast method per line (column E). Use 'Manual' + columns L:O to type amounts directly.", f_sub)
hd = ["Type", "Major Category", "Key", "Description", "Method", "Cash Timing"] + [f"=Inputs!{c}${R_LBL}" for c in YC] + \
     [f'="Manual "&Inputs!{c}${R_LBL}' for c in YC[1:]]
for i, t in enumerate(hd):
    put(ws, f"{L(i+1)}7", t)
hdr_row(ws, 7, 1, 15)
ws.row_dimensions[7].height = 30
put(ws, "D4", "Total enrollment", f_bold); put(ws, "D5", "Tuition index", f_bold); put(ws, "D6", "Salary-method lines total", f_bold)
D1 = 8
DN = D1 + len(lines) - 1
DL = 400
for i, c in enumerate(FY):
    put(ws, f"{c}4", f"=Enrollment!{SC[i]}{E_TOT}", fmt="#,##0")
    put(ws, f"{c}5", f"=Enrollment!{SC[i]}{R_EI}", fmt="0.0000")
    if i == 0:
        put(ws, f"{c}6", f'=SUMIF($E${D1}:$E${DL},"Salary",{c}${D1}:{c}${DL})', fmt=NUM)
    else:  # all 'Salary' lines grow at the same %, so no need to re-sum the column (avoids a circular reference)
        put(ws, f"{c}6", f"={FY[i-1]}6*(1+Inputs!{YC[i]}${R_SAL})", fmt=NUM)
for n, l in enumerate(lines):
    r = D1 + n
    put(ws, f"A{r}", l["type"]); put(ws, f"B{r}", l["cat"])
    put(ws, f"C{r}", l["key"], al=Alignment(horizontal="left")); ws[f"C{r}"].number_format = "@"
    put(ws, f"D{r}", l["desc"])
    inp(ws, f"E{r}", l["method"]); inp(ws, f"F{r}", l["profile"])
    put(ws, f"G{r}", f"=SUMIFS('QB Import'!$B$1:$B${QN},'QB Import'!$H$1:$H${QN},$C{r},'QB Import'!$G$1:$G${QN},1)", f_link, NUM)
    for i in range(1, 5):
        X, P, IC, M = FY[i], FY[i - 1], YC[i], L(11 + i)
        f = (f'=IF($E{r}="Manual",{M}{r},IF($E{r}="Flat",{P}{r},IF($E{r}="Inflation",{P}{r}*(1+Inputs!{IC}${R_INFL}),'
             f'IF($E{r}="Salary",{P}{r}*(1+Inputs!{IC}${R_SAL}),IF($E{r}="Benefits",{P}{r}*(1+Inputs!{IC}${R_BEN}),'
             f'IF($E{r}="Per Student",{P}{r}*{X}$4/{P}$4*(1+Inputs!{IC}${R_INFL}),IF($E{r}="Enrollment",{P}{r}*{X}$4/{P}$4,'
             f'IF($E{r}="Tuition Index",{P}{r}*{X}$5/{P}$5,IF($E{r}="% of Salaries",IF($G$6=0,0,$G{r}/$G$6*{X}$6),'
             f'IF($E{r}="Debt",$G{r}*Inputs!{IC}${R_DEBTMO}/12,IF($E{r}="Reserve Yield",Inputs!{IC}${R_YLD}*Inputs!{IC}${R_INTBASE},{P}{r})))))))))))')
        put(ws, f"{X}{r}", f, fmt=NUM)
        mv = l["manual"][i - 1] if l["manual"] else None
        inp(ws, f"{M}{r}", mv, NUM)
# validations (whole editable block through row 400 so new accounts can be added)
dvm = DataValidation(type="list", formula1='"' + ",".join(METHODS) + '"', allow_blank=True)
dvp = DataValidation(type="list", formula1='"' + ",".join(PROFILES) + '"', allow_blank=True)
ws.add_data_validation(dvm); ws.add_data_validation(dvp)
dvm.add(f"E{D1}:E{DL}"); dvp.add(f"F{D1}:F{DL}")
put(ws, f"A{DN+2}", "To add an account: insert a row above this note (inside the table), fill columns A:F, copy G:K formulas down from the row above.", f_sub)
widths = dict(A=9, B=30, C=13, D=44, E=15, F=13)
for k, v in widths.items():
    ws.column_dimensions[k].width = v
for c in "GHIJKLMNO":
    ws.column_dimensions[c].width = 13
ws.freeze_panes = "E8"
ws.auto_filter.ref = f"A7:O{DN}"

# ================================================================= CASH FLOW
ws = wsC
put(ws, "A1", "MONTHLY CASH FLOW - 5 YEARS", f_title)
put(ws, "A2", "Month ending", f_bold)
put(ws, "A3", "Fiscal year #"); put(ws, "A4", "Fiscal month # (1 = July)"); put(ws, "A5", "Fiscal year", f_bold)
NC = 12 * NY
CF = [L(2 + i) for i in range(NC)]
ws.column_dimensions["A"].width = 44
for c in CF:
    ws.column_dimensions[c].width = 11.5
# CF Summary row map (defined here because Cash Flow reads the profile block)
CS_H = 3
S_BEG, S_REC, S_DIS, S_NET, S_CAP, S_TRF, S_END, S_RES, S_TOT, S_LOW, S_LOWM, S_DAYS = range(4, 16)
SUP_H = 20
SUP_REV1 = SUP_H + 2       # 22..26
SUP_EXP1 = SUP_H + 9       # 29..33
# rows
R_D, R_FY, R_FM, R_FL = 2, 3, 4, 5
R_RECH = 7
RC1 = 8                   # 8..12 receipt rows (5 profiles)
R_RECT = 13
R_DISH = 15
DS1 = 16                  # 16..20
R_STR = 21
R_DIST = 22
R_OPN = 23
R_CAP = 25
R_TRN = 26
R_OBEG, R_ONET, R_OTRF, R_OEND = 28, 29, 30, 31
R_RBEG, R_RIN, R_REND = 33, 34, 35
R_TOTC, R_FLAG = 37, 38
for i, c in enumerate(CF):
    fyi, fm = i // 12 + 1, i % 12 + 1
    put(ws, f"{c}{R_D}", f"=EOMONTH(Inputs!$C${R_START},0)" if i == 0 else f"=EOMONTH({CF[i-1]}{R_D},1)", f_bold, DATE)
    put(ws, f"{c}{R_FY}", fyi); put(ws, f"{c}{R_FM}", fm)
    put(ws, f"{c}{R_FL}", f'="FY"&TEXT(INDEX(Inputs!$C${R_END}:$G${R_END},{c}${R_FY}),"yy")', f_bold)
    ws[f"{c}{R_FL}"].alignment = Alignment(horizontal="right")
hdr_row(ws, R_D, 1, 1 + NC)
for c in CF:
    ws[f"{c}{R_D}"].number_format = DATE
    ws[f"{c}{R_D}"].alignment = Alignment(horizontal="right")
section(ws, R_RECH, "CASH RECEIPTS (by timing profile)", 1 + NC)
section(ws, R_DISH, "CASH DISBURSEMENTS (by timing profile)", 1 + NC)
def prof_formula(c, prow, sup_row):
    pr = f"Inputs!$B${prow}:$O${prow}"
    sup = f"'CF Summary'!$B${sup_row}:$G${sup_row}"
    return (f"=INDEX({sup},{c}${R_FY})*INDEX({pr},{c}${R_FM}+2)"
            f"+IF({c}${R_FM}=11,INDEX({sup},{c}${R_FY}+1)*INDEX({pr},1),0)"
            f"+IF({c}${R_FM}=12,INDEX({sup},{c}${R_FY}+1)*INDEX({pr},2),0)")
for pi, p in enumerate(PROFILES):
    put(ws, f"A{RC1+pi}", f"  {p}")
    put(ws, f"A{DS1+pi}", f"  {p}")
    for c in CF:
        put(ws, f"{c}{RC1+pi}", prof_formula(c, prof_rows[p], SUP_REV1 + pi), fmt=NUM)
        put(ws, f"{c}{DS1+pi}", "=-(" + prof_formula(c, prof_rows[p], SUP_EXP1 + pi)[1:] + ")", fmt=NUM)
put(ws, f"A{R_RECT}", "Total receipts", f_bold)
put(ws, f"A{R_STR}", "  Strategic plan initiatives - operating, net")
put(ws, f"A{R_DIST}", "Total operating disbursements", f_bold)
put(ws, f"A{R_OPN}", "NET OPERATING CASH FLOW", f_bold)
put(ws, f"A{R_CAP}", "Capital projects & capital initiatives")
put(ws, f"A{R_TRN}", "Transfers to reserves")
put(ws, f"A{R_OBEG}", "Operating cash - beginning"); put(ws, f"A{R_ONET}", "  Net operating cash flow"); put(ws, f"A{R_OTRF}", "  Capital & reserve transfers")
put(ws, f"A{R_OEND}", "Operating cash - ending", f_bold)
put(ws, f"A{R_RBEG}", "Reserves & restricted - beginning"); put(ws, f"A{R_RIN}", "  Transfers in"); put(ws, f"A{R_REND}", "Reserves & restricted - ending", f_bold)
put(ws, f"A{R_TOTC}", "TOTAL CASH & RESERVES - ending", f_bold)
put(ws, f"A{R_FLAG}", "Operating cash below zero?", f_sub)
IB = f"Inputs!$B${IT1}:$B${IT2}"
for i, c in enumerate(CF):
    prev = CF[i - 1] if i else None
    put(ws, f"{c}{R_RECT}", f"=SUM({c}{RC1}:{c}{RC1+4})", f_bold, NUM, fill_tot, border=top)
    put(ws, f"{c}{R_STR}", f"=-INDEX(Inputs!$E${R_ITOP}:$I${R_ITOP},{c}${R_FY})/12", fmt=NUM)
    put(ws, f"{c}{R_DIST}", f"=SUM({c}{DS1}:{c}{R_STR})", f_bold, NUM, fill_tot, border=top)
    put(ws, f"{c}{R_OPN}", f"={c}{R_RECT}+{c}{R_DIST}", f_bold, NUM, fill_tot)
    cap_init = (f"SUMPRODUCT(({IB}=\"Capital\")*(Inputs!$C${IT1}:$C${IT2}=\"Y\")*(INDEX(Inputs!$E${IT1}:$I${IT2},0,{c}${R_FY})-INDEX(Inputs!$J${IT1}:$N${IT2},0,{c}${R_FY}))"
                f"*((Inputs!$D${IT1}:$D${IT2}={c}${R_FM})+(Inputs!$D${IT1}:$D${IT2}=0)/12))")
    cap_proj = (f"SUMPRODUCT(INDEX(Inputs!$C${CP1}:$G${CP2},0,{c}${R_FY})*((Inputs!$B${CP1}:$B${CP2}={c}${R_FM})+(Inputs!$B${CP1}:$B${CP2}=0)/12))")
    put(ws, f"{c}{R_CAP}", f"=-({cap_init}+{cap_proj})", fmt=NUM)
    trf = "+".join(f"INDEX(Inputs!$C${r}:$G${r},{c}${R_FY})*((Inputs!$B${r}={c}${R_FM})+(Inputs!$B${r}=0)/12)" for r in (RF1, RF2))
    put(ws, f"{c}{R_TRN}", f"=-({trf})", fmt=NUM)
    put(ws, f"{c}{R_OBEG}", f"=Inputs!$D${R_OBA}" if i == 0 else f"={prev}{R_OEND}", fmt=NUM)
    put(ws, f"{c}{R_ONET}", f"={c}{R_OPN}", fmt=NUM)
    put(ws, f"{c}{R_OTRF}", f"={c}{R_CAP}+{c}{R_TRN}", fmt=NUM)
    put(ws, f"{c}{R_OEND}", f"={c}{R_OBEG}+{c}{R_ONET}+{c}{R_OTRF}", f_bold, NUM, fill_tot, border=top)
    put(ws, f"{c}{R_RBEG}", f"=Inputs!$D${R_OBR}" if i == 0 else f"={prev}{R_REND}", fmt=NUM)
    put(ws, f"{c}{R_RIN}", f"=-{c}{R_TRN}", fmt=NUM)
    put(ws, f"{c}{R_REND}", f"={c}{R_RBEG}+{c}{R_RIN}", f_bold, NUM, fill_tot, border=top)
    put(ws, f"{c}{R_TOTC}", f"={c}{R_OEND}+{c}{R_REND}", f_bold, NUM, fill_tot, border=topbot)
    put(ws, f"{c}{R_FLAG}", f'=IF({c}{R_OEND}<0,"SHORTFALL","")', Font(name=FONT, size=9, bold=True, color="C00000"))
ws.conditional_formatting.add(f"B{R_OEND}:{CF[-1]}{R_OEND}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, bold=True, color="C00000")))
ws.freeze_panes = "B6"

# ================================================================= CF SUMMARY
ws = wsCS
put(ws, "A1", "CASH FLOW SUMMARY (annual)", f_title)
ws.column_dimensions["A"].width = 50
for c in "BCDEFG":
    ws.column_dimensions[c].width = 14
put(ws, f"A{CS_H}", "Fiscal year")
for i, c in enumerate(SC):
    put(ws, f"{c}{CS_H}", f"=Inputs!{YC[i]}${R_LBL}")
hdr_row(ws, CS_H, 1, 6)
labels = {S_BEG: "Operating cash - beginning of year", S_REC: "Cash receipts", S_DIS: "Operating disbursements (incl. strategic plan)",
          S_NET: "Net operating cash flow", S_CAP: "Capital projects & capital initiatives", S_TRF: "Transfers to reserves",
          S_END: "Operating cash - end of year", S_RES: "Reserves & restricted - end of year", S_TOT: "Total cash & reserves - end of year",
          S_LOW: "Lowest month-end operating cash", S_LOWM: "  Month of low point", S_DAYS: "Days cash on hand (total cash / daily expenses)"}
for r, t in labels.items():
    put(ws, f"A{r}", t, f_bold if r in (S_NET, S_END, S_TOT) else f_norm)
CFR = lambda row: f"'Cash Flow'!$B${row}:${CF[-1]}${row}"
def yr_slice(row, i):  # i = 0-based FY index -> INDEX():INDEX() slice
    return f"INDEX({CFR(row)},{i*12+1}):INDEX({CFR(row)},{i*12+12})"
for i, c in enumerate(SC):
    put(ws, f"{c}{S_BEG}", f"=INDEX({CFR(R_OBEG)},{i*12+1})", fmt=NUM)
    put(ws, f"{c}{S_REC}", f"=SUM({yr_slice(R_RECT, i)})", fmt=NUM)
    put(ws, f"{c}{S_DIS}", f"=SUM({yr_slice(R_DIST, i)})", fmt=NUM)
    put(ws, f"{c}{S_NET}", f"={c}{S_REC}+{c}{S_DIS}", f_bold, NUM, fill_tot, border=top)
    put(ws, f"{c}{S_CAP}", f"=SUM({yr_slice(R_CAP, i)})", fmt=NUM)
    put(ws, f"{c}{S_TRF}", f"=SUM({yr_slice(R_TRN, i)})", fmt=NUM)
    put(ws, f"{c}{S_END}", f"=INDEX({CFR(R_OEND)},{i*12+12})", f_bold, NUM, fill_tot, border=top)
    put(ws, f"{c}{S_RES}", f"=INDEX({CFR(R_REND)},{i*12+12})", fmt=NUM)
    put(ws, f"{c}{S_TOT}", f"={c}{S_END}+{c}{S_RES}", f_bold, NUM, fill_tot, border=topbot)
    put(ws, f"{c}{S_LOW}", f"=MIN({yr_slice(R_OEND, i)})", fmt=NUM)
    put(ws, f"{c}{S_LOWM}", f"=INDEX({CFR(R_D)},{i*12}+MATCH({c}{S_LOW},{yr_slice(R_OEND, i)},0))", fmt=DATE, al=Alignment(horizontal="right"))
    put(ws, f"{c}{S_DAYS}", f"=IF(Summary!{c}$24=0,0,{c}{S_TOT}/(Summary!{c}$24/365))", fmt="0")
ws.conditional_formatting.add(f"B{S_LOW}:F{S_LOW}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, bold=True, color="C00000")))
put(ws, f"A{S_DAYS+2}", "Cash flow is on a cash basis: tuition follows the payment-plan profile, prepayments for next year (May/June) are included.", f_sub)
# supporting block
put(ws, f"A{SUP_H-1}", "Supporting calculation: annual amounts by cash-timing profile (feeds the monthly Cash Flow tab)", f_bold)
put(ws, f"A{SUP_H}", "Revenue lines")
for i, c in enumerate(SC):
    put(ws, f"{c}{SUP_H}", f"={c}{CS_H}")
put(ws, f"G{SUP_H}", "Next FY (est.)")
hdr_row(ws, SUP_H, 1, 7)
put(ws, f"A{SUP_H+1}", "  (FY32 is estimated from FY31 x growth; only used for May/June prepayments in the final year)", f_sub)
put(ws, f"A{SUP_EXP1-1}", "Expense lines", f_bold)
for pi, p in enumerate(PROFILES):
    for base_row, typ, gr in ((SUP_REV1, "Revenue", R_TUIT), (SUP_EXP1, "Expense", R_INFL)):
        r = base_row + pi
        put(ws, f"A{r}", f"  {p}")
        for i, c in enumerate(SC):
            put(ws, f"{c}{r}", f"=SUMIFS('Forecast Detail'!{FY[i]}${D1}:{FY[i]}${DL},'Forecast Detail'!$A${D1}:$A${DL},\"{typ}\",'Forecast Detail'!$F${D1}:$F${DL},\"{p}\")", fmt=NUM)
        put(ws, f"G{r}", f"=F{r}*(1+Inputs!$G${gr})", fmt=NUM)
ws.freeze_panes = "B4"

# ================================================================= SUMMARY
ws = wsS
ws.column_dimensions["A"].width = 46
for c in "BCDEF":
    ws.column_dimensions[c].width = 14
ws.column_dimensions["G"].width = 11
put(ws, "A1", "TRINITY PRESBYTERIAN SCHOOL", f_title)
put(ws, "A2", "Five-Year Financial Forecast - Summary by Major Category", Font(name=FONT, size=11, bold=True))
put(ws, "A3", '="Data check: "&B46', f_sub)
put(ws, "A5", "")
for i, c in enumerate(SC):
    put(ws, f"{c}5", f"=Inputs!{YC[i]}${R_LBL}")
put(ws, "G5", "5-yr CAGR")
put(ws, "A5", "Fiscal year ending June 30")
hdr_row(ws, 5, 1, 7)
put(ws, "A6", "Enrollment (students)"); put(ws, "A7", "Tuition rate increase"); put(ws, "A8", "Salary increase")
for i, c in enumerate(SC):
    put(ws, f"{c}6", f"=Enrollment!{c}{E_TOT}", fmt="#,##0")
    put(ws, f"{c}7", f"=Inputs!{YC[i]}{R_TUIT}", fmt=PCT)
    put(ws, f"{c}8", f"=Inputs!{YC[i]}{R_SAL}", fmt=PCT)
rev_cats = ["1 Tuition and Fees", "2 Interest and Other Income", "4 Other Miscellaneous Income", "5 Athletic Income"]
exp_cats = ["1 Salaries and Benefits", "2 Instructional Expenses", "3 Operations", "4 Administrative", "5 Athletic Expenses"]
def cat_sum(c, i, r):
    return (f"=SUMIFS('Forecast Detail'!{FY[i]}${D1}:{FY[i]}${DL},'Forecast Detail'!$B${D1}:$B${DL},$A{r})"
            f"+SUMIFS('QB Import'!$B$1:$B${QN},'QB Import'!$F$1:$F${QN},$A{r},'QB Import'!$I$1:$I${QN},\"UNMAPPED\")")
def cagr(r):
    put(ws, f"G{r}", f'=IF(AND(B{r}>0,F{r}>0),(F{r}/B{r})^(1/4)-1,"")', fmt=PCT)
section(ws, 10, "REVENUE", 7)
r = 11
for cat in rev_cats:
    put(ws, f"A{r}", cat)
    for i, c in enumerate(SC):
        put(ws, f"{c}{r}", cat_sum(c, i, r), fmt=NUM)
    cagr(r); r += 1
R_REV = r   # 15
put(ws, f"A{R_REV}", "Total Revenue", f_bold)
for c in SC:
    put(ws, f"{c}{R_REV}", f"=SUM({c}11:{c}{R_REV-1})", f_bold, NUM, fill_tot, border=top)
cagr(R_REV)
section(ws, 17, "EXPENSES", 7)
r = 18
for cat in exp_cats:
    put(ws, f"A{r}", cat)
    for i, c in enumerate(SC):
        put(ws, f"{c}{r}", cat_sum(c, i, r), fmt=NUM)
    cagr(r); r += 1
R_STRAT = r  # 23
put(ws, f"A{R_STRAT}", "Strategic Plan Initiatives (operating, net)")
for i, c in enumerate(SC):
    put(ws, f"{c}{R_STRAT}", f"=Inputs!{L(5+i)}{R_ITOP}", fmt=NUM)
R_EXP = R_STRAT + 1  # 24
put(ws, f"A{R_EXP}", "Total Expenses", f_bold)
for c in SC:
    put(ws, f"{c}{R_EXP}", f"=SUM({c}18:{c}{R_STRAT})", f_bold, NUM, fill_tot, border=top)
cagr(R_EXP)
R_NOP = 26
put(ws, f"A{R_NOP}", "Net Operating Revenue", f_bold)
put(ws, "A27", "Other Revenue (Development)"); put(ws, "A28", "Other Expenditures (Development)")
put(ws, "A29", "Net Other Revenue"); put(ws, "A30", "NET REVENUE (SURPLUS / DEFICIT)", f_bold)
# hidden lookup keys for other cats
for i, c in enumerate(SC):
    put(ws, f"{c}{R_NOP}", f"={c}{R_REV}-{c}{R_EXP}", f_bold, NUM, fill_tot, border=top)
    for rr, cat in ((27, "Other Revenue"), (28, "Other Expense")):
        put(ws, f"{c}{rr}", f"=SUMIFS('Forecast Detail'!{FY[i]}${D1}:{FY[i]}${DL},'Forecast Detail'!$B${D1}:$B${DL},\"{cat}\")"
            f"+SUMIFS('QB Import'!$B$1:$B${QN},'QB Import'!$F$1:$F${QN},\"{cat}\",'QB Import'!$I$1:$I${QN},\"UNMAPPED\")", fmt=NUM)
    put(ws, f"{c}29", f"={c}27-{c}28", fmt=NUM)
    put(ws, f"{c}30", f"={c}{R_NOP}+{c}29", f_bold, NUM, fill_tot, border=topbot)
section(ws, 32, "KEY METRICS", 7)
mets = [(33, "Salaries & benefits % of total revenue", lambda c: f"=IF({c}{R_REV}=0,0,{c}18/{c}{R_REV})", PCT),
        (34, "Net tuition & fees per student", lambda c: f"=IF({c}6=0,0,{c}11/{c}6)", NUM),
        (35, "Net revenue margin (% of revenue)", lambda c: f"=IF({c}{R_REV}=0,0,{c}30/{c}{R_REV})", PCT),
        (36, "Capital projects & initiatives (cash)", lambda c: f"=-'CF Summary'!{c}{S_CAP}", NUM),
        (37, "Ending cash & reserves", lambda c: f"='CF Summary'!{c}{S_TOT}", NUM),
        (38, "Lowest month-end operating cash", lambda c: f"='CF Summary'!{c}{S_LOW}", NUM),
        (39, "Days cash on hand", lambda c: f"='CF Summary'!{c}{S_DAYS}", "0")]
for rr, lab, fn, fmt in mets:
    put(ws, f"A{rr}", lab)
    for c in SC:
        put(ws, f"{c}{rr}", fn(c), fmt=fmt)
ws.conditional_formatting.add("B30:F30", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, bold=True, color="C00000")))
ws.conditional_formatting.add("B38:F38", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=FONT, color="C00000")))
put(ws, "A41", "Source: FY27 = QuickBooks budget import; FY28-FY31 = forecast per Inputs and Forecast Detail. Cash on a cash basis.", f_sub)
# ---- checks (outside print area)
put(ws, "A43", "IMPORT CHECKS (not printed)", f_bold)
Q = lambda col: f"'QB Import'!${col}$1:${col}${QN}"
chk = [(44, "Revenue: import detail vs summary FY27",
        f'=SUMIFS({Q("B")},{Q("E")},"Revenue",{Q("G")},1)+SUMIFS({Q("B")},{Q("E")},"COGS",{Q("G")},1)-SUM(B11:B{R_REV-1})'),
       (45, "Expenses: import detail vs summary FY27",
        f'=SUMIFS({Q("B")},{Q("E")},"Expense",{Q("G")},1)-SUM(B18:B{R_STRAT-1})'),
       (47, "Other net: import detail vs summary FY27",
        f'=SUMIFS({Q("B")},{Q("E")},"Other Revenue",{Q("G")},1)-SUMIFS({Q("B")},{Q("E")},"Other Expense",{Q("G")},1)-B29'),
       (48, "Net Revenue: summary FY27 vs QuickBooks 'Net Revenue' line",
        f'=IF(COUNTIF({Q("D")},"Net Revenue")=0,0,SUMIFS({Q("B")},{Q("D")},"Net Revenue")-B30)'),
       (49, "Accounts on import not in Forecast Detail (count)", f'=COUNTIF({Q("I")},"UNMAPPED")'),
       (50, "Forecast Detail rows with no method (count)", f"=COUNTA('Forecast Detail'!$C${D1}:$C${DL})-COUNTA('Forecast Detail'!$E${D1}:$E${DL})")]
for rr, lab, f in chk:
    put(ws, f"A{rr}", lab); put(ws, f"B{rr}", f, fmt=NUM)
put(ws, "A46", "Overall status", f_bold)
put(ws, "B46", '=IF(AND(ABS(B44)<1,ABS(B45)<1,ABS(B47)<1,ABS(B48)<1,B49=0,B50=0),"OK - ties to QuickBooks","CHECK - "&B49&" unmapped account(s); see rows 44-50")', f_bold)
ws.conditional_formatting.add("B46", FormulaRule(formula=['LEFT(B46,2)="OK"'], font=Font(name=FONT, bold=True, color="00803C")))
ws.conditional_formatting.add("B46", FormulaRule(formula=['LEFT(B46,5)="CHECK"'], font=Font(name=FONT, bold=True, color="C00000")))
ws.print_area = "A1:G41"
ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
ws.print_options.horizontalCentered = True
ws.freeze_panes = "B6"

for w, ori in ((wsCS, "landscape"), (wsC, "landscape"), (wsF, "landscape"), (wsI, "landscape"), (wsE, "portrait")):
    w.page_setup.orientation = ori
    w.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
    w.page_setup.fitToWidth = 1
    w.page_setup.fitToHeight = 0
wsC.print_title_cols = "A:A"
wsC.page_setup.fitToWidth = 0

# ================================================================= READ ME
ws = wsR
ws.column_dimensions["A"].width = 120
txt = ["FIVE-YEAR FORECAST MODEL - READ ME", "",
       "TABS",
       "  Summary          One-page print view by QuickBooks Major Category (the accounts under '1 Tuition and Fees', '1 Salaries and Benefits', etc.).",
       "  Inputs           All assumptions: growth %, strategic plan initiatives, capital projects, reserve funding, opening cash, payment plans, cash timing.",
       "  Enrollment       FY27 enrollment and rates by grade; later years roll forward by grade using retention and new-student inputs.",
       "  Forecast Detail  One row per QuickBooks account. Pick a forecast method and cash timing per line.",
       "  Cash Flow        Monthly cash flow, 60 months (5 fiscal years, July-June).",
       "  CF Summary       Annual cash flow summary, low-point cash and days cash on hand.",
       "  QB Import        Paste the QuickBooks budget export here (A1). Helper formulas in D:I map every account to its Major Category.", "",
       "HOW MAJOR CATEGORIES ARE FOUND",
       "  Any row whose label begins with a single digit and a space (e.g. '1 Tuition and Fees') starts a Major Category. Every account beneath it, including",
       "  subgroups such as 'Gross Tuition' or 'Tuition Discounts', rolls up into it until the next such row or the next section (Revenue / Expenditures / Other Revenue / Other Expenditures).",
       "  Rows that carry an amount but have no account number (e.g. 'Food Service') are treated as their own detail line.", "",
       "UPDATING FROM A NEW QUICKBOOKS EXPORT",
       "  1. Clear QB Import columns A:C, paste the new export at A1.  2. Read the status on Summary!A3.  3. If accounts show as UNMAPPED, add them to Forecast Detail.",
       "  Unmapped accounts are still counted (held flat) so totals never drop out of the forecast.", "",
       "FORECAST METHODS (Forecast Detail column E)",
       "  Flat              Same as prior year.",
       "  Inflation         Prior year x (1 + general inflation).",
       "  Salary            Prior year x (1 + salary increase).",
       "  Benefits          Prior year x (1 + benefits increase).",
       "  Per Student       Prior year x enrollment change x (1 + inflation).",
       "  Enrollment        Prior year x enrollment change.",
       "  Tuition Index     Prior year x change in the tuition index (enrollment by grade x rate by grade). Calibrated to the QuickBooks FY27 budget.",
       "  % of Salaries     Keeps the FY27 ratio to total 'Salary' lines (payroll taxes, 401K).",
       "  Debt              FY27 amount x months of payments remaining (see final mortgage payment date).",
       "  Reserve Yield     Yield % x prior year-end interest-earning balances (Inputs).",
       "  Manual            Type the amount for each year in Forecast Detail columns L:O.", "",
       "SIMPLIFICATIONS TO KNOW ABOUT",
       "  - Mortgage principal and interest are kept as expenses exactly as in QuickBooks; interest is not stepped down as principal is repaid.",
       "  - Interest income is based on prior year-end balances (opening balances + reserve funding), not the monthly cash flow, to avoid circularity.",
       "  - The cash flow starts July of FY27 from the opening balances on Inputs; May/June collections for FY27 tuition are assumed to be in those balances.",
       "  - Scholarship programs (S4K, Choose Act, AOSF) are driven by the method chosen on their own lines, not by student counts.",
       "  - Strategic Plan Initiatives are incremental to the QuickBooks account '500265 Strategic Plan Expenses'."]
for i, t in enumerate(txt, start=1):
    put(ws, f"A{i}", t, f_title if i == 1 else (f_bold if t.isupper() or t.startswith(("TABS", "HOW", "UPDATING", "FORECAST", "SIMPLIF")) else f_norm))

wb.save(OUT)
print("saved", OUT, "lines:", len(lines))
